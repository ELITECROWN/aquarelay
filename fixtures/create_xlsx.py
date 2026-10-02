"""Reproduce the synthetic XLSX fixture without macros or formulas."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

root = Path(__file__).resolve().parent
rows = json.loads((root / "demo-ngo-missing-unit.json").read_text())
book = Workbook()
sheet = book.active
sheet.title = "Synthetic demo measurements"
columns = list(rows[0])
sheet.append(columns)
for row in rows:
    sheet.append([row[column] for column in columns])
for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="183B30")
for column, width in {"A": 22, "B": 29, "C": 27, "D": 16, "E": 16, "F": 15}.items():
    sheet.column_dimensions[column].width = width
sheet.freeze_panes = "A2"
book.properties.title = "AquaRelay synthetic NGO import fixture"
book.properties.description = "Synthetic values and fictional water-body names. Units and timezone intentionally omitted for manual mapping demonstration. CC0."
book.save(root / "demo-ngo-missing-unit.xlsx")
