"""Reviewed mappings, durable import runs, safe connectors and local handoffs.

Remote sources are data. No source text is executed or passed to an AI model.
"""
import csv
import hashlib
import hmac
import http.client
import io
import ipaddress
import json
import math
import os
import socket
import ssl
import time
from datetime import datetime, timezone, timedelta
from pathlib import Path
from urllib.parse import urlsplit, urljoin, parse_qsl
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response, FileResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.orm import Session

from .db import get_db
from .auth import require_manager
from .models import Connector, MappingVersion, ImportRun, Receipt, WaterBody, Observation, Source, Job, Audit, User, uid, utcnow
from .core import emit_event
from .standards import build_fhir, check_fhir, build_sensorthings, sensorthings_rows

router = APIRouter(tags=["integrations"])
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100
DESTINATIONS = {"site_name", "waterbody_id", "observed_at", "source_updated_at", "temperature", "ph", "dissolved_oxygen", "external_id", "parameter", "value", "unit", "latitude", "longitude", "photo"}
PARAMETERS = {"temperature", "ph", "dissolved_oxygen"}


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str).encode()).hexdigest()


def schema_hash(rows):
    return digest(sorted({str(k) for row in rows for k in row}))


def csv_safe(value):
    text = str(value if value is not None else "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def parse_upload(filename, payload):
    if len(payload) > MAX_BYTES:
        raise ValueError("File exceeds the 5 MiB limit")
    extension = Path(filename or "").suffix.lower()
    if extension == ".csv":
        try:
            reader = csv.DictReader(io.StringIO(payload.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            if not headers or any(not str(header).strip() for header in headers) or len(set(headers)) != len(headers):
                raise ValueError("CSV requires unique nonempty column headers")
            rows = list(reader)
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("CSV must be valid UTF-8 with a header") from exc
    elif extension == ".json":
        try:
            decoded = json.loads(payload)
        except (ValueError, UnicodeError) as exc:
            raise ValueError("Invalid JSON") from exc
        rows = decoded if isinstance(decoded, list) else decoded.get("records") if isinstance(decoded, dict) else None
        if not isinstance(rows, list):
            raise ValueError("JSON must contain an array of records, or {records: [...]}")
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("Expanded spreadsheet exceeds the 40 MiB limit")
                if any("vbaProject" in item.filename for item in archive.infolist()):
                    raise ValueError("Macros are unsupported")
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
            sheet = book.active
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValueError("Spreadsheet exceeds row/column limits")
            values = []
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Formulas are unsupported; export literal values first")
                values.append([cell.value.isoformat() if isinstance(cell.value, datetime) else cell.value for cell in cells])
            book.close()
            if not values:
                raise ValueError("Spreadsheet is empty")
            headers = [str(x or "").strip() for x in values[0]]
            if not all(headers) or len(set(headers)) != len(headers):
                raise ValueError("Spreadsheet requires unique nonempty column headers")
            rows = [dict(zip(headers, row)) for row in values[1:] if any(x is not None for x in row)]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Invalid XLSX workbook") from exc
    else:
        raise ValueError("Supported files: .csv, .xlsx and .json. Macros are never accepted.")
    if not rows or len(rows) > MAX_ROWS:
        raise ValueError("Provide 1–2000 source records")
    if any(not isinstance(row, dict) or len(row) > MAX_COLUMNS or None in row for row in rows):
        raise ValueError("Each record requires at most 100 named fields")
    if any(len(str(value)) > 16000 for row in rows for value in row.values()):
        raise ValueError("A source cell exceeds the 16,000-character limit")
    return rows


def suggest_mapping(columns):
    aliases = {"site_name": "site_name", "waterbody_id": "waterbody_id", "observed_at": "observed_at", "water_temp": "temperature", "temperature": "temperature",
               "ph_level": "ph", "ph": "ph", "external_id": "external_id", "event_id": "external_id", "unit": "unit", "parameter": "parameter", "value": "value",
               "source_updated_at": "source_updated_at", "latitude": "latitude", "longitude": "longitude", "photo": "photo"}
    aliases.update({'site':'site_name','location':'site_name','date':'observed_at','time':'observed_at','timestamp':'observed_at','temp':'temperature','watertemperature':'temperature','wt':'temperature','ph_lvl':'ph','phlevel':'ph','do':'dissolved_oxygen','dissolvedoxygen':'dissolved_oxygen'})
    return {column: aliases[column.lower().strip().replace(' ','_')] for column in columns if column.lower().strip().replace(' ','_') in aliases}


def parse_time(value, timezone_name):
    if not value:
        raise ValueError("missing_time")
    try:
        moment = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (ValueError, TypeError) as exc:
        raise ValueError("invalid_time") from exc
    if moment.tzinfo is None:
        if not timezone_name:
            raise ValueError("missing_timezone")
        try:
            zone = ZoneInfo(timezone_name)
        except ZoneInfoNotFoundError as exc:
            raise ValueError("invalid_timezone") from exc
        earlier, later = moment.replace(tzinfo=zone, fold=0), moment.replace(tzinfo=zone, fold=1)
        if earlier.utcoffset() != later.utcoffset():
            raise ValueError("ambiguous_timezone")
        if earlier.astimezone(timezone.utc).astimezone(zone).replace(tzinfo=None) != moment:
            raise ValueError("nonexistent_time")
        moment = earlier
    return moment.astimezone(timezone.utc).isoformat()


def transform_rows(rows, mapping, units, timezone_name, waterbodies, waterbody_id=None, parameter_map=None):
    if any(destination not in DESTINATIONS for destination in mapping.values() if destination):
        raise ValueError("Unsupported mapping destination")
    destinations = [destination for destination in mapping.values() if destination]
    if len(destinations) != len(set(destinations)):
        raise ValueError("A destination may be mapped only once")
    results = []
    for index, raw in enumerate(rows):
        mapped = {destination: raw.get(source) for source, destination in mapping.items() if destination}
        issues, observations = [], []
        def issue(code, message, fatal=False):
            issues.append({"row": index + 1, "code": code, "message": message, "severity": "rejected" if fatal else "unresolved"})
        site = waterbody_id or mapped.get("waterbody_id") or mapped.get("site_name")
        matches = [body for body in waterbodies if str(site or "").strip().casefold() in {str(body.id).casefold(), body.name.casefold(), *(str(x).casefold() for x in (body.aliases or []))}]
        body = matches[0] if len(matches) == 1 else None
        if not body:
            issue("ambiguous_site" if len(matches) > 1 else "unknown_site", "Select a unique water-body identity explicitly.")
        try:
            observed = parse_time(mapped.get("observed_at"), timezone_name)
        except ValueError as exc:
            observed = None
            code = str(exc)
            issue(code, "Observation time requires an explicit timezone or UTC offset." if code == "missing_timezone" else "Resolve the invalid, missing or ambiguous observation time.", code in {"invalid_time", "missing_time"})
        source_updated = None
        if mapped.get("source_updated_at"):
            try:
                source_updated = parse_time(mapped["source_updated_at"], timezone_name)
            except ValueError as exc:
                issue("source_updated_" + str(exc), "Source update time must be a valid explicit timestamp.", str(exc) == "invalid_time")
        if mapped.get("external_id") and len(str(mapped["external_id"])) > 200:
            issue("invalid_external_id", "Source event IDs must be at most 200 characters.", True)
        for key, limit in [("latitude", 90), ("longitude", 180)]:
            if mapped.get(key) not in (None, ""):
                try:
                    number = float(mapped[key])
                    if not math.isfinite(number) or abs(number) > limit:
                        raise ValueError()
                except (ValueError, TypeError):
                    issue("invalid_coordinate", "Coordinates must be finite decimal degrees within valid ranges.", True)
        measurements = [(key, mapped[key]) for key in PARAMETERS if mapped.get(key) not in (None, "")]
        if mapped.get("value") not in (None, ""):
            parameter = str(mapped.get("parameter") or "")
            parameter = (parameter_map or {}).get(parameter, parameter)
            if parameter not in PARAMETERS:
                issue("unknown_parameter", "Explicitly map the source parameter; abbreviations are not guessed.")
            else:
                measurements.append((parameter, mapped["value"]))
        if not measurements and not any(x["code"] == "unknown_parameter" for x in issues):
            issue("missing_measurement", "Map at least one supported environmental measurement.", True)
        for parameter, raw_value in measurements:
            source_unit = units.get(parameter) or (mapped.get("unit") if len(measurements) == 1 else None)
            if not source_unit:
                issue("missing_unit", "Choose the source unit for " + parameter + "; none is assumed.")
                continue
            unit_aliases = {"degC": "Cel", "°C": "Cel", "C": "Cel", "celsius": "Cel", "Cel": "Cel", "degF": "[degF]", "°F": "[degF]", "F": "[degF]", "[degF]": "[degF]", "1": "1", "mg/L": "mg/L"}
            canonical_unit = unit_aliases.get(str(source_unit))
            allowed = {"temperature": {"Cel", "[degF]"}, "ph": {"1"}, "dissolved_oxygen": {"mg/L"}}[parameter]
            if canonical_unit not in allowed:
                issue("unsupported_unit", "Explicit supported units: Celsius/Fahrenheit for temperature, 1 for pH, mg/L for dissolved oxygen.")
                continue
            try:
                value = float(raw_value)
                if not math.isfinite(value) or isinstance(raw_value, bool):
                    raise ValueError()
            except (ValueError, TypeError):
                issue("invalid_value", "Measurement must be a finite number.", True)
                continue
            conversion = "Source unit retained"
            if canonical_unit == "[degF]":
                value = (value - 32) * 5 / 9
                canonical_unit = "Cel"
                conversion = "Approved Fahrenheit to Celsius: (source value − 32) × 5/9"
            observations.append({"waterbody_id": body.id if body else None, "parameter": parameter, "value": value, "unit": canonical_unit,
                  "observed_at": observed, "external_id": str(mapped.get("external_id") or digest(raw)), "source_updated_at": source_updated,
                  "conversion": conversion, "source_value": raw_value, "source_unit": source_unit})
        results.append({"row": index + 1, "raw": raw, "observations": observations, "issues": issues,
                        "status": "rejected" if any(x["severity"] == "rejected" for x in issues) else "unresolved" if issues else "ready"})
    return results


def validate_public_url(url):
    """Resolve only public HTTPS:443. Reject mixed public/private DNS responses."""
    try:
        parsed = urlsplit(url)
        if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password or parsed.port not in (None, 443) or parsed.fragment:
            raise ValueError("Connector URL must be public HTTPS on port 443, without credentials or fragment")
        credential_keys = {"key", "api_key", "apikey", "token", "access_token", "access_key", "client_secret", "auth", "secret", "password", "credential", "signature", "authorization"}
        if any(key.lower().replace("-", "_") in credential_keys for key, value in parse_qsl(parsed.query)):
            raise ValueError("Credential query parameters are unsupported. Store credentials in protected headers instead.")
        hostname = parsed.hostname.encode("idna").decode("ascii")
        addresses = sorted({item[4][0] for item in socket.getaddrinfo(hostname, 443, type=socket.SOCK_STREAM)})
    except (OSError, UnicodeError) as exc:
        raise ValueError("Connector host could not be resolved") from exc
    if not addresses or any(not ipaddress.ip_address(address).is_global or ipaddress.ip_address(address).is_multicast for address in addresses):
        raise ValueError("Connector resolved to a nonpublic address; private and reserved destinations are blocked")
    return parsed, addresses


class PinnedHTTPSConnection(http.client.HTTPSConnection):
    def __init__(self, hostname, address):
        super().__init__(hostname, port=443, timeout=8, context=ssl.create_default_context())
        self.public_address = address

    def connect(self):
        # Connect exactly to the checked address, retaining TLS SNI and certificate checks.
        sock = socket.create_connection((self.public_address, 443), self.timeout)
        self.sock = self._context.wrap_socket(sock, server_hostname=self.host)


def safe_fetch_json(url, headers=None):
    initial_host = urlsplit(url).hostname
    for redirect in range(4):
        parsed, addresses = validate_public_url(url)
        connection = PinnedHTTPSConnection(parsed.hostname, addresses[0])
        request_headers = {"Accept": "application/json", "User-Agent": "AquaRelay/1.0"}
        if parsed.hostname == initial_host:
            request_headers.update(headers or {})
        try:
            target = parsed.path or "/"
            if parsed.query:
                target += "?" + parsed.query
            connection.request("GET", target, headers=request_headers)
            response = connection.getresponse()
            if response.status in {301, 302, 303, 307, 308}:
                location = response.getheader("Location")
                if not location or redirect == 3:
                    raise ValueError("Redirect limit exceeded or missing target")
                url = urljoin(url, location)
                continue
            if response.status < 200 or response.status >= 300:
                raise ValueError("Remote HTTP " + str(response.status))
            data = response.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise ValueError("Remote response exceeds 5 MiB")
            try:
                return json.loads(data)
            except (ValueError, UnicodeError) as exc:
                raise ValueError("Remote response is not valid JSON") from exc
        finally:
            connection.close()
    raise ValueError("Redirect limit exceeded")


def verify_webhook(secret, body, timestamp, signature, now=None):
    try:
        if abs((now if now is not None else time.time()) - int(timestamp)) > 300:
            return False
        expected = hmac.new(secret.encode(), str(timestamp).encode() + b"." + body, hashlib.sha256).hexdigest()
        return hmac.compare_digest(expected, str(signature).removeprefix("sha256="))
    except (ValueError, TypeError, AttributeError):
        return False


def parse_recipient_response(status_code, payload, idempotency_key):
    """HTTP acceptance is delivery; acknowledgement requires an explicit bound receipt."""
    result = {"status_code": status_code, "receipt_id": None, "acknowledged": False}
    if not isinstance(payload, dict):
        return result
    remote_id = payload.get("receipt_id")
    if isinstance(remote_id, str) and 0 < len(remote_id) <= 160:
        result["receipt_id"] = remote_id
    acknowledgement = payload.get("acknowledgement")
    if payload.get("contract") != "aquarelay-handoff-v1" or payload.get("idempotency_key") != idempotency_key or not result["receipt_id"] or not isinstance(acknowledgement, dict):
        return result
    ack_id = acknowledgement.get("receipt_id")
    if acknowledgement.get("status") != "acknowledged" or not isinstance(ack_id, str) or not 0 < len(ack_id) <= 160:
        return result
    try:
        received_at = parse_time(acknowledgement.get("received_at"), None)
    except ValueError:
        return result
    result.update(acknowledged=True, acknowledgement={"receipt_id": ack_id, "received_at": received_at,
        "status": "acknowledged", "contract": "aquarelay-handoff-v1"})
    return result


def safe_post_fhir(url, payload, headers, idempotency_key):
    """Bounded configured FHIR POST. Redirects retain POST only within the same origin."""
    initial_host = urlsplit(url).hostname
    encoded = json.dumps(payload, allow_nan=False, separators=(",", ":")).encode()
    if len(encoded) > MAX_BYTES:
        raise ValueError("Outgoing FHIR payload exceeds 5 MiB")
    for redirect in range(4):
        parsed, addresses = validate_public_url(url)
        if parsed.hostname != initial_host:
            raise ValueError("Institutional POST redirects may not change recipient hostname")
        connection = PinnedHTTPSConnection(parsed.hostname, addresses[0])
        request_headers = {"Content-Type": "application/fhir+json", "Accept": "application/json", "User-Agent": "AquaRelay/1.0", **headers,
                           "Idempotency-Key": idempotency_key}
        try:
            target = parsed.path or "/"
            if parsed.query:
                target += "?" + parsed.query
            connection.request("POST", target, body=encoded, headers=request_headers)
            response = connection.getresponse()
            if response.status in {307, 308}:
                location = response.getheader("Location")
                if not location or redirect == 3:
                    raise ValueError("Institutional POST redirect limit exceeded")
                url = urljoin(url, location)
                continue
            if not 200 <= response.status < 300:
                raise ValueError("Institutional endpoint returned HTTP " + str(response.status))
            body = response.read(MAX_BYTES + 1)
            if len(body) > MAX_BYTES:
                raise ValueError("Institutional response exceeds 5 MiB")
            try:
                decoded = json.loads(body) if body else {}
            except (ValueError, UnicodeError):
                decoded = {}
            result = parse_recipient_response(response.status, decoded, idempotency_key)
            header_id = response.getheader("X-Receipt-ID")
            if not result["receipt_id"] and isinstance(header_id, str) and 0 < len(header_id) <= 160:
                result["receipt_id"] = header_id
            return result
        finally:
            connection.close()
    raise ValueError("Institutional POST redirect limit exceeded")


def institutional_recipient(organisation_id, check_url=False):
    """Recipients are configured by server operators, never from a public fetch field."""
    try:
        configured = json.loads(os.getenv("HANDOFF_RECIPIENTS_JSON", "{}"))
        config = configured.get(organisation_id) if isinstance(configured, dict) else None
        if not isinstance(config, dict):
            raise ValueError("No institutional recipient is configured for your organisation. Export remains available.")
        if config.get("contract") != "aquarelay-handoff-v1" or config.get("idempotency_supported") is not True:
            raise ValueError("The institutional recipient must declare the supported idempotent handoff contract.")
        if not isinstance(config.get("name"), str) or not 0 < len(config["name"]) <= 160 or not isinstance(config.get("url"), str):
            raise ValueError("Institutional recipient configuration is incomplete.")
        headers = config.get("headers", {})
        if not isinstance(headers, dict) or any(name.lower() not in {"authorization", "x-api-key"} or not isinstance(value, str) or "\r" in value or "\n" in value for name, value in headers.items()):
            raise ValueError("Institutional credential configuration is unsupported.")
        if check_url:
            validate_public_url(config["url"])
        return config
    except (ValueError, TypeError, AttributeError) as exc:
        explicit_reason = isinstance(exc, ValueError) and ("configuration" in str(exc).lower() or str(exc).startswith(("No institutional", "The institutional")))
        reason = str(exc) if explicit_reason else "Institutional recipient configuration is invalid or its destination is blocked."
        raise HTTPException(503, reason) from exc


class MappingInput(BaseModel):
    mapping: dict[str, str]
    units: dict[str, str] = Field(default_factory=dict)
    timezone: str | None = None
    waterbody_id: str | None = None
    parameter_map: dict[str, str] = Field(default_factory=dict)


class ConnectorInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    kind: str = "http_json"
    url: str | None = None


class ConfigureInput(BaseModel):
    url: str | None = None
    headers: dict[str, str] | None = None
    webhook_secret: str | None = None
    mapping: dict[str, str] | None = None
    units: dict[str, str] | None = None
    timezone: str | None = None
    waterbody_id: str | None = None
    parameter_map: dict[str, str] | None = None
    stale_after_hours: int | None = Field(default=None, ge=1, le=8760)
    poll_interval_minutes: int | None = Field(default=None, ge=1, le=10080)


def owned(db, model, record_id, user):
    record = db.get(model, record_id)
    if not record:
        raise HTTPException(404, "Record not found")
    if record.organisation_id != user.organisation_id:
        raise HTTPException(403, "This record belongs to another organisation")
    return record


def connector_dict(record):
    config = record.config or {}
    state = record.state
    if state == "healthy" and record.last_success_at:
        now = datetime.now(timezone.utc)
        threshold = timedelta(hours=config.get("stale_after_hours", 72))
        if now - datetime.fromisoformat(record.last_success_at) > threshold or record.last_observed_at and now - datetime.fromisoformat(record.last_observed_at) > threshold:
            state = "stale"
    return {"id": record.id, "name": record.name, "kind": record.kind, "state": state, "status": state,
            "url": config.get("url"), "configured": bool(config.get("url") or record.kind in {"file", "demo_fixture"}),
            "credentials_configured": bool(config.get("headers") or config.get("webhook_secret")),
            "mapping": config.get("mapping", {}), "units": config.get("units", {}), "timezone": config.get("timezone"),
            "waterbody_id": config.get("waterbody_id"), "parameter_map": config.get("parameter_map", {}),
            "mapping_approved": bool(config.get("approved_mapping_version")), "poll_interval_minutes": config.get("poll_interval_minutes"),
            "polling_enabled": bool(config.get("poll_interval_minutes") and config.get("approved_mapping_version") and record.state not in {"paused", "disabled"}),
            "next_poll_at": config.get("next_poll_at"), "demo_scenario": config.get("fixture_state") if record.kind == "demo_fixture" else None,
            "last_attempt_at": record.last_attempt_at,
            "last_success_at": record.last_success_at, "last_observed_at": record.last_observed_at, "error": record.error,
            "synthetic": bool((record.data or {}).get("synthetic")), **{k: v for k, v in (record.data or {}).items() if k in {"imported", "quarantined", "retry_history", "next_retry_at", "schema_changed"}}}


def result_counts(records):
    return {"ready": sum(x["status"] == "ready" for x in records), "imported": sum(x["status"] == "imported" for x in records),
            "duplicate": sum(x["status"] == "duplicate" for x in records), "rejected": sum(x["status"] == "rejected" for x in records),
            "unresolved": sum(x["status"] == "unresolved" for x in records), "total": len(records)}


def run_dict(run, include_raw=False):
    data = run.data or {}
    result = {"id": run.id, "connector_id": run.connector_id, "state": run.state, "created_at": run.created_at,
            "approved_at": run.approved_at, "filename": data.get("filename"), "synthetic": data.get("synthetic", False),
            "mapping_version_id": run.mapping_version_id, "source_id": run.source_id, "mapping_settings": data.get("settings", {}),
            "schema_changed": data.get("schema_changed", False), "counts": data.get("counts", {}), "error": data.get("error")}
    result.update(data.get("counts", {}))
    if include_raw:
        result.update({"columns": data.get("columns", []), "rows": data.get("rows", []), "suggestions": data.get("suggestions", {}),
                       "records": data.get("records", []), "issues": [issue for row in data.get("records", []) for issue in row["issues"]]})
    return result


def create_run(db, user, rows, filename, connector=None, synthetic=False):
    if len(rows) > MAX_ROWS or not rows or any(not isinstance(row, dict) or len(row) > MAX_COLUMNS for row in rows):
        raise HTTPException(422, "Source requires 1–2000 object records with at most 100 fields")
    columns = sorted({str(k) for row in rows for k in row})
    fingerprint = schema_hash(rows)
    latest = db.scalar(select(MappingVersion).where(MappingVersion.connector_id == connector.id).order_by(MappingVersion.version.desc())) if connector else None
    run = ImportRun(id=uid("imp"), connector_id=connector.id if connector else None, organisation_id=user.organisation_id, source_hash=digest(rows),
          state="preview", data={"filename": filename, "rows": rows, "columns": columns, "source_hash": digest(rows),
          "schema_hash": fingerprint, "schema_changed": bool(latest and latest.schema_hash != fingerprint),
          "suggestions": suggest_mapping(columns), "synthetic": synthetic}, created_at=utcnow())
    db.add(run)
    db.flush()
    return run


def transform_run(db, run, settings):
    records = transform_rows(run.data["rows"], settings["mapping"], settings.get("units", {}), settings.get("timezone"), db.scalars(select(WaterBody)).all(), settings.get("waterbody_id"), settings.get("parameter_map"))
    synthetic_places = {body.id for body in db.scalars(select(WaterBody).where(WaterBody.synthetic.is_(True))).all()}
    synthetic = run.data.get("synthetic", False) or any(o["waterbody_id"] in synthetic_places for row in records for o in row["observations"])
    run.data = {**run.data, "settings": settings, "records": records, "counts": result_counts(records), "synthetic": synthetic}
    run.state = "transformed"
    return run


def approve_run(db, run, user):
    if run.state == "approved":
        return run
    if run.state != "transformed":
        raise HTTPException(409, "Preview the transformed records before approving this import")
    connector = db.get(Connector, run.connector_id) if run.connector_id else None
    if not connector:
        # Identical organisation/dataset label shares an identity across reuploads.
        source_key = "file:" + user.organisation_id + ":" + run.data["filename"]
        connector_id = "con-file-" + digest(source_key)[:24]
        connector = db.get(Connector, connector_id)
        if not connector:
            connector = Connector(id=connector_id, organisation_id=user.organisation_id, name=run.data["filename"], kind="file", state="configured", config={}, data={"synthetic": run.data["synthetic"]}, created_at=utcnow(), updated_at=utcnow())
            db.add(connector)
            db.flush()
        run.connector_id = connector.id
    source_id = "src-import-" + digest(connector.id)[:24]
    source = db.get(Source, source_id)
    if not source:
        source = Source(id=source_id, waterbody_id=None, organisation_id=user.organisation_id, connector_id=connector.id, name=connector.name,
                        kind="import", synthetic=run.data["synthetic"], data={"connector_id": connector.id, "origin": "Approved source import"}, created_at=utcnow())
        db.add(source)
        db.flush()
    settings = run.data["settings"]
    version_no = (db.scalar(select(func.max(MappingVersion.version)).where(MappingVersion.connector_id == connector.id)) or 0) + 1
    version = MappingVersion(id=uid("map"), connector_id=connector.id, organisation_id=user.organisation_id, version=version_no, schema_hash=run.data["schema_hash"], mapping=settings["mapping"], data={"settings": settings, "approved_by": user.id}, created_at=utcnow())
    db.add(version)
    db.flush()
    run.mapping_version_id = version.id
    run.source_id = source.id
    records = json.loads(json.dumps(run.data["records"]))
    affected = set()
    for record in records:
        if record["status"] != "ready":
            continue
        existing = []
        for measurement in record["observations"]:
            row = db.scalar(select(Observation).where(Observation.source_id == source_id, Observation.external_id == measurement["external_id"], Observation.parameter == measurement["parameter"]))
            if row:
                if row.waterbody_id != measurement["waterbody_id"] or row.value != measurement["value"] or row.unit != measurement["unit"] or row.observed_at != measurement["observed_at"]:
                    record["issues"].append({"row": record["row"], "code": "source_correction", "message": "Source event already exists with different values. Correction review is required; the prior record is preserved.", "severity": "unresolved"})
                existing.append(measurement["parameter"])
        if record["issues"]:
            record["status"] = "unresolved"
            continue
        if len(existing) == len(record["observations"]):
            record["status"] = "duplicate"
            continue
        for measurement in record["observations"]:
            if measurement["parameter"] in existing:
                continue
            source_updated = None
            if measurement.get("source_updated_at"):
                try:
                    source_updated = parse_time(measurement["source_updated_at"], settings.get("timezone"))
                except ValueError:
                    pass
            sta = record["raw"].get("sensorthings") or {}
            supplied_party = sta.get("party") or {} if isinstance(sta, dict) else {}
            supplied_license = sta.get("license") or {} if isinstance(sta, dict) else {}
            public_party = {key: supplied_party[key] for key in {"role", "displayName", "description"} if isinstance(supplied_party, dict) and key in supplied_party}
            public_license = {key: supplied_license[key] for key in {"name", "description", "definition", "attributionText"} if isinstance(supplied_license, dict) and key in supplied_license}
            db.add(Observation(id=uid("obs"), waterbody_id=measurement["waterbody_id"], source_id=source_id,
                 external_id=measurement["external_id"], parameter=measurement["parameter"], value=measurement["value"], unit=measurement["unit"],
                 observed_at=measurement["observed_at"], received_at=utcnow(), source_updated_at=source_updated,
                 synthetic=run.data["synthetic"], data={"raw": record["raw"], "raw_payload_reference": run.id, "mapping_version_id": version.id,
                 "source_value": measurement["source_value"], "source_unit": measurement["source_unit"], "conversion": measurement["conversion"],
                 "staplus_party": public_party, "staplus_license": public_license}))
            affected.add(measurement["waterbody_id"])
        record["status"] = "imported"
        db.flush()
    counts = result_counts(records)
    run.state = "approved"
    run.approved_at = utcnow()
    run.data = {**run.data, "records": records, "counts": counts, "approved_by": user.id}
    connector.config = {**(connector.config or {}), **settings, "approved_mapping_version": version.id, "schema_hash": run.data["schema_hash"]}
    connector.state = "healthy"
    connector.last_success_at = utcnow()
    connector.error = None
    times = [measurement["observed_at"] for record in records if record["status"] in {"imported", "duplicate"} for measurement in record["observations"]]
    if times:
        connector.last_observed_at = max(times)
    connector.data = {**(connector.data or {}), "imported": (connector.data or {}).get("imported", 0) + counts["imported"],
                      "quarantined": counts["unresolved"] + counts["rejected"], "schema_changed": False, "synthetic": run.data["synthetic"]}
    source.observed_at = connector.last_observed_at
    source.state = connector_dict(connector)["state"]
    source.received_at = utcnow()
    updated_times = [measurement["source_updated_at"] for record in records if record["status"] in {"imported", "duplicate"} for measurement in record["observations"] if measurement.get("source_updated_at")]
    if updated_times:
        source.source_updated_at = max(updated_times)
    source.synthetic = run.data["synthetic"]
    db.add(Audit(id=uid("audit"), actor_id=user.id, kind="import_approved", target_id=run.id, data={"mapping_version_id": version.id, "counts": counts}))
    for body_id in affected:
        emit_event(db, body_id, None, "source_import", "Source records imported", f"{counts['imported']} source rows approved. Source values and mapping decisions are preserved.", user.id, source_id=source.id)
    db.flush()
    return run


@router.get("/connectors")
def list_connectors(db: Session = Depends(get_db), user=Depends(require_manager)):
    return {"items": [connector_dict(x) for x in db.scalars(select(Connector).where(Connector.organisation_id == user.organisation_id)).all()]}


@router.post("/connectors", status_code=201)
def add_connector(payload: ConnectorInput, db: Session = Depends(get_db), user=Depends(require_manager)):
    if payload.kind not in {"http_json", "sensorthings", "file", "demo_fixture"}:
        raise HTTPException(422, "Supported connector kinds: http_json, sensorthings, file, demo_fixture")
    if payload.kind == "demo_fixture" and os.getenv("DEMO_MODE", "true").lower() != "true":
        raise HTTPException(404, "Demo scenario sources are disabled outside DEMO_MODE")
    if payload.url:
        try:
            validate_public_url(payload.url)
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
    row = Connector(id=uid("con"), organisation_id=user.organisation_id, name=payload.name, kind=payload.kind,
          state="configured" if payload.url or payload.kind == "demo_fixture" else "disabled", config={"url": payload.url, **({"fixture_state": "healthy"} if payload.kind == "demo_fixture" else {})},
          data={"synthetic": payload.kind == "demo_fixture"}, created_at=utcnow(), updated_at=utcnow())
    db.add(row)
    db.commit()
    return connector_dict(row)


@router.get("/connectors/{connector_id}")
def connector_detail(connector_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    runs = db.scalars(select(ImportRun).where(ImportRun.connector_id == connector_id).order_by(ImportRun.created_at.desc()).limit(30)).all()
    versions = db.scalars(select(MappingVersion).where(MappingVersion.connector_id == connector_id).order_by(MappingVersion.version.desc())).all()
    return {"connector": connector_dict(connector), "runs": [run_dict(x) for x in runs], "mapping_versions": [{"id": x.id, "version": x.version, "schema_hash": x.schema_hash, "mapping": x.mapping, "created_at": x.created_at} for x in versions]}


@router.post("/connectors/{connector_id}/configure")
def configure(connector_id: str, payload: ConfigureInput, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    incoming = payload.model_dump(exclude_unset=True)
    if incoming.get("url"):
        try:
            validate_public_url(incoming["url"])
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
    for name, value in (incoming.get("headers") or {}).items():
        if name.lower() not in {"authorization", "x-api-key"} or "\r" in value or "\n" in value:
            raise HTTPException(422, "Only Authorization or X-API-Key credentials are supported")
    if any(key in incoming for key in {"mapping", "units", "timezone", "waterbody_id", "parameter_map"}):
        incoming["approved_mapping_version"] = None
    connector.config = {**(connector.config or {}), **incoming}
    connector.state = "configured" if connector.config.get("url") or connector.kind in {"file", "demo_fixture"} else "disabled"
    connector.updated_at = utcnow()
    connector.error = None
    db.commit()
    return connector_dict(connector)


@router.post("/imports/preview", status_code=201)
async def preview(file: UploadFile = File(...), connector_id: str | None = Form(None), label: str | None = Form(None), synthetic: bool = Form(False), db: Session = Depends(get_db), user=Depends(require_manager)):
    payload = await file.read(MAX_BYTES + 1)
    try:
        rows = parse_upload(file.filename, payload)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    connector = owned(db, Connector, connector_id, user) if connector_id else None
    synthetic = synthetic or all(row.get("synthetic") is True or str(row.get("synthetic", "")).lower() == "true" for row in rows)
    run = create_run(db, user, rows, label or Path(file.filename or "source").name, connector, synthetic)
    db.commit()
    return run_dict(run, True)


@router.get("/imports/{run_id}")
def import_detail(run_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    return run_dict(owned(db, ImportRun, run_id, user), True)


@router.get("/imports/sample/{format}")
def sample_import(format: str, user=Depends(require_manager)):
    if format not in {"csv", "json", "xlsx"}:
        raise HTTPException(404, "Supported sample formats: csv, json, xlsx")
    filename = "demo-ngo-missing-unit." + format
    path = Path(__file__).resolve().parents[2] / "fixtures" / filename
    if not path.exists():
        raise HTTPException(503, "Sample fixture files are unavailable in this installation")
    media = {"csv": "text/csv", "json": "application/json", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[format]
    return FileResponse(path, media_type=media, filename=filename, headers={"X-AquaRelay-Data-Origin": "synthetic-demo-fixture"})


@router.post("/imports/{run_id}/transform")
def transform(run_id: str, payload: MappingInput, db: Session = Depends(get_db), user=Depends(require_manager)):
    run = owned(db, ImportRun, run_id, user)
    if run.state == "approved":
        raise HTTPException(409, "Approved import is immutable; upload a new run for correction review")
    try:
        transform_run(db, run, payload.model_dump())
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    db.commit()
    return run_dict(run, True)


@router.post("/imports/{run_id}/approve")
def approve(run_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    run = owned(db, ImportRun, run_id, user)
    approve_run(db, run, user)
    db.commit()
    return run_dict(run, True)


def record_failure(db, connector, message, kind):
    # Never expose exception text containing credentials, full URLs or source values.
    safe_message = message if isinstance(message, str) and len(message) < 300 and not any(value in message for value in (connector.config or {}).get("headers", {}).values()) else "Connector request failed"
    connector.state = "error"
    connector.error = safe_message
    data = connector.data or {}
    history = list(data.get("retry_history", []))
    attempt = len(history) + 1
    next_retry = datetime.now(timezone.utc) + timedelta(seconds=min(3600, 15 * 2 ** min(attempt - 1, 8)))
    history.append({"attempt": attempt, "at": utcnow(), "kind": kind, "error": safe_message, "next_retry_at": next_retry.isoformat()})
    connector.data = {**data, "retry_history": history[-50:], "next_retry_at": next_retry.isoformat()}
    db.commit()


def remote_rows(connector):
    if connector.kind == "demo_fixture":
        if os.getenv("DEMO_MODE", "true").lower() != "true":
            raise ValueError("Demo scenario source is disabled outside DEMO_MODE")
        if (connector.config or {}).get("fixture_state") == "failure":
            raise ValueError("Demo scenario: synthetic source is deliberately unavailable")
        return [
            {"external_id": "demo-source-temperature-001", "site_name": "Demo Reedwater Lake", "observed_at": "2026-09-30T06:00:00Z", "parameter": "temperature", "value": 24.8, "unit": "Cel", "synthetic": True},
            {"external_id": "demo-source-ph-001", "site_name": "Demo Reedwater Lake", "observed_at": "2026-09-30T06:00:00Z", "parameter": "ph", "value": 7.2, "unit": "1", "synthetic": True}
        ]
    if connector.kind == "file":
        raise ValueError("File sources are updated through Upload and review")
    url = (connector.config or {}).get("url")
    if not url:
        raise ValueError("Connector is disabled until a public HTTPS endpoint is configured")
    payload = safe_fetch_json(url, (connector.config or {}).get("headers", {}))
    if connector.kind == "sensorthings":
        return sensorthings_rows(payload)
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict) and isinstance(payload.get("records"), list):
        return payload["records"]
    raise ValueError("HTTP JSON requires an array of records or {records: [...]}.")


@router.post("/connectors/{connector_id}/test")
def test_connector(connector_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    connector.last_attempt_at = utcnow()
    db.commit()
    try:
        rows = remote_rows(connector)
        run = create_run(db, user, rows, connector.name, connector, bool((connector.data or {}).get("synthetic")))
        connector.state = "configured"
        connector.error = None
        db.commit()
        return {"ok": True, "message": "Source fetched. Review mapping and approve the preview before syncing.", "connector": connector_dict(connector), "preview": run_dict(run, True), "run_id": run.id}
    except (ValueError, OSError, http.client.HTTPException, HTTPException) as exc:
        record_failure(db, connector, str(exc.detail) if isinstance(exc, HTTPException) else str(exc), "test")
        return {"ok": False, "error": connector.error, "connector": connector_dict(connector)}


@router.post("/connectors/{connector_id}/pause")
def pause_connector(connector_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    connector.state = "paused"
    connector.updated_at = utcnow()
    db.commit()
    return connector_dict(connector)


class ScenarioInput(BaseModel):
    state: str


@router.post("/connectors/{connector_id}/scenario")
def demo_scenario(connector_id: str, payload: ScenarioInput, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    if connector.kind != "demo_fixture" or os.getenv("DEMO_MODE", "true").lower() != "true":
        raise HTTPException(404, "This is available only for an explicitly labelled local demo scenario source")
    if payload.state not in {"failure", "recovered"}:
        raise HTTPException(422, "Demo scenario state must be failure or recovered")
    connector.config = {**(connector.config or {}), "fixture_state": "failure" if payload.state == "failure" else "healthy"}
    if payload.state == "recovered":
        # Explicit simulation reset; live remote connectors retain their retry backoff.
        connector.data = {**(connector.data or {}), "next_retry_at": None}
    db.add(Audit(id=uid("audit"), actor_id=user.id, target_id=connector.id, kind="demo_source_scenario", data={"synthetic": True, "state": payload.state}))
    db.commit()
    return {"message": "Demo scenario: " + payload.state + ". Run Test or Sync to record the resulting attempt.", "connector": connector_dict(connector)}


@router.post("/connectors/{connector_id}/sync")
def sync_connector(connector_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    connector = owned(db, Connector, connector_id, user)
    config = connector.config or {}
    if connector.state == "paused":
        raise HTTPException(409, "Resume the connector by configuring it before syncing")
    if not config.get("approved_mapping_version"):
        raise HTTPException(409, "Test the source, review its transformed preview and approve a mapping before syncing")
    if (connector.data or {}).get("next_retry_at") and datetime.fromisoformat(connector.data["next_retry_at"]) > datetime.now(timezone.utc):
        raise HTTPException(429, "Retry backoff is active", headers={"Retry-After": "15"})
    connector.last_attempt_at = utcnow()
    db.commit()
    try:
        rows = remote_rows(connector)
        run = create_run(db, user, rows, connector.name, connector, bool((connector.data or {}).get("synthetic")))
        if run.data["schema_changed"]:
            connector.state = "error"
            connector.error = "Source schema changed. Review and approve a new mapping version."
            connector.data = {**(connector.data or {}), "schema_changed": True}
            db.commit()
            return {"ok": False, "error": connector.error, "preview": run_dict(run, True), "run_id": run.id, "connector": connector_dict(connector)}
        version = db.get(MappingVersion, config["approved_mapping_version"])
        transform_run(db, run, version.data["settings"])
        approve_run(db, run, user)
        connector.data = {**(connector.data or {}), "next_retry_at": None}
        db.commit()
        return {"ok": True, "connector": connector_dict(connector), "run": run_dict(run, True), **run.data["counts"]}
    except (ValueError, OSError, http.client.HTTPException, HTTPException) as exc:
        db.rollback()
        connector = db.get(Connector, connector_id)
        record_failure(db, connector, str(exc.detail) if isinstance(exc, HTTPException) else str(exc), "sync")
        return {"ok": False, "error": connector.error, "connector": connector_dict(connector)}


def schedule_due_connectors(db):
    """Queue one durable polling job per due source, never before approval."""
    now = utcnow()
    eligible = select(Connector).where(Connector.state.notin_(["paused", "disabled"]))
    if db.bind.dialect.name == "postgresql":
        eligible = eligible.with_for_update(skip_locked=True)
    for connector in db.scalars(eligible).all():
        config = connector.config or {}
        interval = config.get("poll_interval_minutes")
        if not interval or not config.get("approved_mapping_version") or connector.kind == "file":
            continue
        if connector.kind == "demo_fixture" and os.getenv("DEMO_MODE", "true").lower() != "true":
            continue
        if config.get("next_poll_at") and config["next_poll_at"] > now:
            continue
        pending = db.scalar(select(Job.id).where(Job.kind == "connector_sync", Job.dedup_key.like("connector_sync:" + connector.id + ":%"), Job.state.in_(["pending", "running"])))
        if pending:
            continue
        next_retry = (connector.data or {}).get("next_retry_at")
        available = max(now, next_retry) if next_retry else now
        key = "connector_sync:" + connector.id + ":" + (config.get("next_poll_at") or connector.created_at)
        if not db.scalar(select(Job.id).where(Job.dedup_key == key)):
            db.add(Job(id=uid("job"), kind="connector_sync", dedup_key=key, state="pending", available_at=available, data={"connector_id": connector.id}))
        connector.config = {**config, "next_poll_at": (datetime.now(timezone.utc) + timedelta(minutes=interval)).isoformat()}
    db.commit()


def process_connector_job(db, job):
    connector = db.get(Connector, job.data["connector_id"])
    if not connector or connector.state in {"paused", "disabled"}:
        return
    config = connector.config or {}
    version = db.get(MappingVersion, config.get("approved_mapping_version")) if config.get("approved_mapping_version") else None
    if not version:
        return
    user = db.get(User, version.data.get("approved_by"))
    if not user or user.organisation_id != connector.organisation_id:
        raise ValueError("Connector mapping approver is unavailable")
    require_manager(user, db)
    result = sync_connector(connector.id, db=db, user=user)
    if not result.get("ok"):
        raise ValueError("Connector polling attempt failed; see retained integration history")


@router.post("/webhooks/{connector_id}")
async def webhook(connector_id: str, request: Request, db: Session = Depends(get_db)):
    connector = db.get(Connector, connector_id)
    body = await request.body()
    if len(body) > MAX_BYTES:
        raise HTTPException(413, "Webhook exceeds 5 MiB")
    secret = (connector.config or {}).get("webhook_secret") if connector else None
    if not secret or not verify_webhook(secret, body, request.headers.get("X-AquaRelay-Timestamp"), request.headers.get("X-AquaRelay-Signature")):
        raise HTTPException(401, "Invalid or expired webhook signature")
    if connector.state == "paused":
        raise HTTPException(409, "Connector paused")
    try:
        payload = json.loads(body)
        external_id = str(payload["event_id"])
        rows = payload["records"]
    except (ValueError, KeyError, TypeError):
        raise HTTPException(422, "Webhook requires event_id and records")
    receipt_id = "rec-webhook-" + digest([connector_id, external_id])[:32]
    previous = db.get(Receipt, receipt_id)
    if previous:
        if previous.data.get("payload_hash") != hashlib.sha256(body).hexdigest():
            raise HTTPException(409, "Event ID was already received with a different payload")
        return {"receipt_id": previous.id, "replay": True, "run_id": previous.data["run_id"], "state": previous.state}
    from types import SimpleNamespace
    user = SimpleNamespace(id="webhook:" + connector.id, organisation_id=connector.organisation_id)
    run = create_run(db, user, rows, connector.name, connector, bool((connector.data or {}).get("synthetic")))
    # Webhook receipts queue a reviewable preview. They never bypass mapping approval.
    receipt = Receipt(id=receipt_id, waterbody_id=None, connector_id=connector.id, organisation_id=connector.organisation_id,
                      kind="webhook", state="received", data={"event_id": external_id, "payload_hash": hashlib.sha256(body).hexdigest(), "run_id": run.id}, created_at=utcnow())
    db.add(receipt)
    db.commit()
    return {"receipt_id": receipt.id, "replay": False, "run_id": run.id, "state": "received", "message": "Source retained for mapping review; no authoritative observations imported yet."}


def export_records(db, waterbody_id):
    waterbody = db.get(WaterBody, waterbody_id)
    if not waterbody:
        raise HTTPException(404, "Water body not found")
    records = db.scalars(select(Observation).where(Observation.waterbody_id == waterbody_id).order_by(Observation.observed_at.desc()).limit(2000)).all()
    return waterbody, records


@router.get("/standards/{waterbody_id}/fhir")
def fhir_export(waterbody_id: str, db: Session = Depends(get_db)):
    payload = build_fhir(*export_records(db, waterbody_id))
    return Response(json.dumps(payload), media_type="application/fhir+json", headers={"X-AquaRelay-Validation": "structural-subset-only; FHIR R4 4.0.1; no full validator"})


@router.get("/standards/{waterbody_id}/sensorthings")
def sensor_export(waterbody_id: str, db: Session = Depends(get_db)):
    return build_sensorthings(*export_records(db, waterbody_id))


class HandoffInput(BaseModel):
    format: str = "fhir"
    client_id: str | None = Field(default=None, max_length=120)
    recipient: str | None = None


def receipt_dict(receipt):
    return {"id": receipt.id, "waterbody_id": receipt.waterbody_id, "kind": receipt.kind, "state": receipt.state,
            "status": receipt.state, "created_at": receipt.created_at, **{key: value for key, value in (receipt.data or {}).items() if key not in {"payload"}}}


@router.get("/handoff-capabilities")
def handoff_capabilities(user=Depends(require_manager)):
    try:
        target = institutional_recipient(user.organisation_id, check_url=True)
        institutional = {"available": True, "name": target["name"], "reason": None}
    except HTTPException as exc:
        institutional = {"available": False, "name": None, "reason": exc.detail}
    return {"local_demo": {"available": os.getenv("DEMO_MODE", "true").lower() == "true", "name": "AquaRelay local demo receiver"},
            "institutional": {**institutional, "supported_formats": ["fhir"], "acknowledgement_contract": "aquarelay-handoff-v1"}}


@router.post("/handoffs/{waterbody_id}", status_code=202)
def handoff(waterbody_id: str, payload: HandoffInput = HandoffInput(), db: Session = Depends(get_db), user=Depends(require_manager)):
    channel = payload.recipient or ("local_demo" if os.getenv("DEMO_MODE", "true").lower() == "true" else "institutional")
    if channel not in {"local_demo", "institutional"}:
        raise HTTPException(422, "Supported recipients: local_demo or institutional")
    target = institutional_recipient(user.organisation_id, check_url=True) if channel == "institutional" else None
    if channel == "local_demo" and os.getenv("DEMO_MODE", "true").lower() != "true":
        raise HTTPException(503, "Local demo receiver disabled. Export remains available.")
    if channel == "institutional" and payload.format != "fhir":
        raise HTTPException(422, "The configured institutional adapter supports FHIR R4 Bundles only")
    if payload.format not in {"fhir", "sensorthings"}:
        raise HTTPException(422, "Supported formats: fhir or sensorthings")
    body, records = export_records(db, waterbody_id)
    exported = build_fhir(body, records) if payload.format == "fhir" else build_sensorthings(body, records)
    check = check_fhir(exported) if payload.format == "fhir" else {"valid": True, "level": "structural-subset", "full_validation": False}
    if not check["valid"]:
        raise HTTPException(422, {"message": "Payload structural checks failed", "validation": check})
    receipt_id = "rec-handoff-" + digest([user.organisation_id, waterbody_id, payload.client_id])[:32] if payload.client_id else uid("rec")
    existing = db.get(Receipt, receipt_id)
    if existing:
        if existing.data.get("channel", "local_demo") != channel or existing.data["format"] != payload.format:
            raise HTTPException(409, "This client ID already identifies a different handoff channel or format")
        return receipt_dict(existing)
    recipient_name = target["name"] if target else "AquaRelay local demo receiver"
    receipt = Receipt(id=receipt_id, waterbody_id=waterbody_id, connector_id=None, organisation_id=user.organisation_id,
                     kind="institutional_handoff" if target else "local_demo_handoff", state="queued",
                     data={"format": payload.format, "payload": exported, "payload_hash": digest(exported), "validation": check,
                     "synthetic": body.synthetic, "recipient": recipient_name, "channel": channel, "acknowledged": False,
                     "recipient_configuration_id": digest([target["url"], target["name"], target["contract"]]) if target else None,
                     "message": "Queued for durable delivery. This does not mean an organisation acknowledged the record."}, created_at=utcnow())
    db.add(receipt)
    db.flush()
    db.add(Job(id=uid("job"), kind="handoff", dedup_key="handoff:" + receipt.id, state="pending", data={"receipt_id": receipt.id}, created_at=utcnow()))
    db.commit()
    return receipt_dict(receipt)


@router.get("/receipts")
def receipts(db: Session = Depends(get_db), user=Depends(require_manager)):
    rows = db.scalars(select(Receipt).where(Receipt.organisation_id == user.organisation_id).order_by(Receipt.created_at.desc()).limit(100)).all()
    return {"items": [receipt_dict(row) for row in rows]}


@router.get("/receipts/{receipt_id}")
def receipt_detail(receipt_id: str, db: Session = Depends(get_db), user=Depends(require_manager)):
    row = owned(db, Receipt, receipt_id, user)
    return {**receipt_dict(row), "payload": (row.data or {}).get("payload")}


def accept_demo_payload(db, payload, organisation_id, waterbody_id=None, external_id=None):
    if os.getenv("DEMO_MODE", "true").lower() != "true":
        raise HTTPException(404, "Local demo receiver disabled")
    if payload.get("resourceType") == "Bundle":
        validation = check_fhir(payload)
    else:
        try:
            rows = sensorthings_rows(payload)
            validation = {"valid": bool(rows), "level": "structural-subset", "full_validation": False}
        except ValueError as exc:
            validation = {"valid": False, "level": "structural-subset", "errors": [str(exc)], "full_validation": False}
    if not validation["valid"]:
        raise HTTPException(422, {"message": "Unsupported payload or structural errors", "validation": validation})
    receipt_id = "rec-receiver-" + digest([organisation_id, external_id or digest(payload)])[:32]
    previous = db.get(Receipt, receipt_id)
    if previous:
        return previous
    receipt = Receipt(id=receipt_id, organisation_id=organisation_id, waterbody_id=waterbody_id, connector_id=None,
            kind="local_demo_receiver", state="received", created_at=utcnow(), data={"payload": payload, "payload_hash": digest(payload), "validation": validation,
            "synthetic": "synthetic" in json.dumps(payload).lower(), "acknowledged": False, "recipient": "AquaRelay local demo receiver",
            "message": "Local demo endpoint accepted and stored this payload. No municipality, hospital or organisation acknowledgement is claimed."})
    db.add(receipt)
    db.flush()
    return receipt


@router.post("/demo/receiver", status_code=201)
async def demo_receiver(request: Request, db: Session = Depends(get_db), user=Depends(require_manager)):
    body = await request.body()
    if len(body) > MAX_BYTES:
        raise HTTPException(413, "Payload exceeds 5 MiB")
    try:
        payload = json.loads(body)
        if not isinstance(payload, dict):
            raise ValueError()
    except ValueError:
        raise HTTPException(422, "Expected JSON object")
    receipt = accept_demo_payload(db, payload, user.organisation_id)
    db.commit()
    return receipt_dict(receipt)


def process_handoff_job(db, job):
    receipt = db.get(Receipt, job.data["receipt_id"])
    if not receipt or receipt.state == "delivered":
        return
    if receipt.data.get("channel") == "institutional":
        try:
            target = institutional_recipient(receipt.organisation_id, check_url=True)
            if receipt.data["recipient_configuration_id"] != digest([target["url"], target["name"], target["contract"]]):
                raise ValueError("Institutional recipient changed after the handoff was queued")
            result = safe_post_fhir(target["url"], receipt.data["payload"], target.get("headers", {}), receipt.id)
            received_id = "rec-external-" + digest(receipt.id)[:32]
            received = db.get(Receipt, received_id)
            external_key = "external:" + digest([receipt.organisation_id, receipt.data["recipient_configuration_id"], result.get("receipt_id")]) if result.get("receipt_id") else None
            if not received and external_key:
                received = db.scalar(select(Receipt).where(Receipt.external_id == external_key))
            if not received:
                received = Receipt(id=received_id, waterbody_id=receipt.waterbody_id, organisation_id=receipt.organisation_id, connector_id=None,
                    kind="institutional_delivery", state="received", created_at=utcnow(),
                    external_id=external_key,
                    data={"recipient": receipt.data["recipient"], "channel": "institutional", "external_receipt_id": result.get("receipt_id"),
                          "http_status": result["status_code"], "synthetic": receipt.data.get("synthetic", False), "acknowledged": result.get("acknowledged", False),
                          "acknowledgement": result.get("acknowledgement"), "validation": receipt.data["validation"]})
                db.add(received)
                db.flush()
            receipt.state = "delivered"
            receipt.data = {**receipt.data, "received_receipt_id": received.id, "external_receipt_id": result.get("receipt_id"),
                "delivered_at": utcnow(), "http_status": result["status_code"], "acknowledged": result.get("acknowledged", False),
                "acknowledgement": result.get("acknowledgement"), "error": None,
                "message": "Configured recipient returned an explicit acknowledgement receipt." if result.get("acknowledged") else "Configured endpoint accepted the payload. Organisation acknowledgement has not been received."}
            return
        except (ValueError, OSError, http.client.HTTPException, HTTPException):
            attempts = job.attempts or 1
            history = list(receipt.data.get("retry_history", []))
            history.append({"attempt": attempts, "at": utcnow(), "error": "Institutional delivery attempt failed; no acknowledgement received."})
            receipt.state = "failed" if attempts >= 5 else "retrying"
            receipt.data = {**receipt.data, "error": "Institutional delivery attempt failed. Check server recipient configuration or endpoint availability.",
                            "retry_history": history[-50:], "acknowledged": False, "last_attempt_at": utcnow()}
            db.commit()
            raise ValueError("Institutional handoff attempt failed")
    received = accept_demo_payload(db, receipt.data["payload"], receipt.organisation_id, receipt.waterbody_id, receipt.id)
    receipt.state = "delivered"
    receipt.data = {**receipt.data, "received_receipt_id": received.id, "delivered_at": utcnow(), "acknowledged": False,
                    "message": "Delivered to the local demo receiver and stored. Organisation acknowledgement has not been received."}
